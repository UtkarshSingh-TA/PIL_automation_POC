import os
import io
import re
import csv
import email
from pathlib import Path
from typing import List, Dict, Tuple
from email.utils import parsedate_to_datetime

from pypdf import PdfReader
import pandas as pd
import docx
import extract_msg
from bs4 import BeautifulSoup


class FileAttachmentExtractor:
    """Helper methods to extract text from PDF, Excel, Word, and CSV attachments."""

    @staticmethod
    def extract_text_from_pdf_bytes(pdf_bytes: bytes) -> str:
        try:
            reader = PdfReader(io.BytesIO(pdf_bytes))
            extracted_pages = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    extracted_pages.append(text)
            return "\n".join(extracted_pages)
        except Exception as e:
            print(f"  └─► [PDF Extractor Warning]: {e}")
            return ""

    @staticmethod
    def extract_text_from_excel_bytes(excel_bytes: bytes, filename: str) -> str:
        try:
            excel_file = pd.ExcelFile(io.BytesIO(excel_bytes))
            sheet_texts = []
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                sheet_str = df.to_string(index=False)
                sheet_texts.append(f"--- Sheet: {sheet_name} ---\n{sheet_str}")
            return "\n".join(sheet_texts)
        except Exception as e:
            print(f"  └─► [Excel Extractor Warning] ({filename}): {e}")
            return ""

    @staticmethod
    def extract_text_from_docx_bytes(docx_bytes: bytes, filename: str) -> str:
        try:
            doc = docx.Document(io.BytesIO(docx_bytes))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text.append(" | ".join(row_text))
            return "\n".join(full_text)
        except Exception as e:
            print(f"  └─► [Word Extractor Warning] ({filename}): {e}")
            return ""

    @staticmethod
    def extract_text_from_csv_bytes(csv_bytes: bytes, filename: str) -> str:
        try:
            content = csv_bytes.decode("utf-8", errors="ignore")
            reader = csv.reader(io.StringIO(content))
            rows = [" | ".join(row) for row in reader if row]
            return "\n".join(rows)
        except Exception as e:
            print(f"  └─► [CSV Extractor Warning] ({filename}): {e}")
            return ""

    @classmethod
    def process_attachment(cls, filename: str, file_bytes: bytes) -> str:
        """Routes attachments to appropriate parser based on file extension."""
        fname_lower = filename.lower()
        if fname_lower.endswith(".pdf"):
            return cls.extract_text_from_pdf_bytes(file_bytes)
        elif fname_lower.endswith((".xlsx", ".xls")):
            return cls.extract_text_from_excel_bytes(file_bytes, filename)
        elif fname_lower.endswith((".docx", ".doc")):
            return cls.extract_text_from_docx_bytes(file_bytes, filename)
        elif fname_lower.endswith(".csv"):
            return cls.extract_text_from_csv_bytes(file_bytes, filename)
        return ""


class OutlookFileParser:
    """Parses text body, dates, and attachments (PDF, Excel, Word, CSV) from .eml and .msg files."""

    @staticmethod
    def clean_html_body(raw_text_or_html: str) -> str:
        """Strips HTML tags, cleans &nbsp; entities, removes null bytes, and normalizes line breaks."""
        if not raw_text_or_html:
            return ""
        
        raw_text_or_html = raw_text_or_html.replace("\x00", "")

        if "<div" in raw_text_or_html or "<html" in raw_text_or_html or "<br" in raw_text_or_html or "<p" in raw_text_or_html:
            soup = BeautifulSoup(raw_text_or_html, "html.parser")
            text = soup.get_text(separator="\n")
        else:
            text = raw_text_or_html

        text = text.replace("\xa0", " ").replace("&nbsp;", " ")
        text = re.sub(r"\n\s*\n", "\n", text)
        return text.strip()

    @staticmethod
    def split_and_reverse_thread(full_text: str) -> str:
        """
        Splits an email thread by standard Outlook reply headers (EN, ES, DE, FR)
        and reverses the blocks so the original request appears first (bottom-to-top order).
        """
        if not full_text:
            return ""

        reply_pattern = r"(?i)(\n-{3,}\s*Original Message\s*-{3,}|\nFrom:\s+|\nDe:\s+|\nEnviado el:\s+|\nVon:\s+)"
        thread_parts = re.split(reply_pattern, full_text)
        
        if len(thread_parts) <= 1:
            return full_text

        reconstructed_messages = []
        reconstructed_messages.append(thread_parts[0]) 
        
        for i in range(1, len(thread_parts), 2):
            header = thread_parts[i]
            body = thread_parts[i+1] if i+1 < len(thread_parts) else ""
            reconstructed_messages.append(header + body)

        # Reverse order: Oldest message (bottom) -> Newest reply (top)
        chronological_thread = list(reversed(reconstructed_messages))
        
        structured_payload = []
        for index, msg in enumerate(chronological_thread, start=1):
            if msg.strip():
                structured_payload.append(f"--- [THREAD STEP #{index} (CHRONOLOGICAL ORDER)] ---\n{msg.strip()}")
            
        return "\n\n".join(structured_payload)

    @classmethod
    def parse_eml_file(cls, file_path: Path) -> Tuple[str, str, str, str, List[str]]:
        """Parses standard MIME .eml files."""
        try:
            with open(file_path, "rb") as f:
                msg = email.message_from_binary_file(f)

            subject = msg.get("Subject", "No Subject")
            sender = msg.get("From", "Unknown Sender")
            raw_date = msg.get("Date", "")
            
            try:
                msg_date = str(parsedate_to_datetime(raw_date)) if raw_date else "Unknown Date"
            except Exception:
                msg_date = raw_date or "Unknown Date"

            body_text = ""
            attachment_texts = []
            filenames = []

            for part in msg.walk():
                content_type = part.get_content_type()
                content_disposition = str(part.get("Content-Disposition"))
                filename = part.get_filename()

                if content_type == "text/plain" and "attachment" not in content_disposition:
                    payload = part.get_payload(decode=True)
                    if payload:
                        body_text += payload.decode("utf-8", errors="ignore")

                elif content_type == "text/html" and not body_text and "attachment" not in content_disposition:
                    payload = part.get_payload(decode=True)
                    if payload:
                        body_text += payload.decode("utf-8", errors="ignore")

                elif "attachment" in content_disposition and filename:
                    file_bytes = part.get_payload(decode=True)
                    if file_bytes:
                        extracted_text = FileAttachmentExtractor.process_attachment(filename, file_bytes)
                        if extracted_text.strip():
                            attachment_texts.append(f"--- ATTACHMENT CONTENT ({filename}) ---\n{extracted_text}")
                            filenames.append(filename)

            cleaned_body = cls.clean_html_body(body_text)
            chronological_body = cls.split_and_reverse_thread(cleaned_body)

            full_text = chronological_body
            if attachment_texts:
                full_text += "\n\n" + "\n\n".join(attachment_texts)

            return subject, sender, msg_date, full_text, filenames

        except Exception as e:
            print(f"  └─► [EML Parsing Error] Failed to parse '{file_path.name}': {e}")
            return f"Error reading {file_path.name}", "Unknown Sender", "Unknown Date", "", []

    @classmethod
    def parse_msg_file(cls, file_path: Path) -> Tuple[str, str, str, str, List[str]]:
        """Parses Outlook binary .msg files using extract_msg."""
        try:
            msg = extract_msg.Message(str(file_path))
            
            subject = str(msg.subject or "No Subject")
            sender = str(msg.sender or "Unknown Sender")
            msg_date = str(msg.date or "Unknown Date")
            
            raw_body = msg.body or ""
            if not raw_body.strip() and hasattr(msg, 'htmlBody') and msg.htmlBody:
                raw_body = msg.htmlBody.decode('utf-8', errors='ignore') if isinstance(msg.htmlBody, bytes) else str(msg.htmlBody)

            body_text = cls.clean_html_body(raw_body)
            chronological_body = cls.split_and_reverse_thread(body_text)

            attachment_texts = []
            filenames = []

            if hasattr(msg, 'attachments') and msg.attachments:
                for att in msg.attachments:
                    try:
                        filename = getattr(att, 'longFilename', None) or getattr(att, 'shortFilename', None) or "attachment"
                        file_bytes = att.data
                        if file_bytes:
                            extracted_text = FileAttachmentExtractor.process_attachment(str(filename), file_bytes)
                            if extracted_text.strip():
                                attachment_texts.append(f"--- ATTACHMENT CONTENT ({filename}) ---\n{extracted_text}")
                                filenames.append(str(filename))
                    except Exception as att_err:
                        print(f"  └─► [MSG Attachment Warning] Skipping attachment in {file_path.name}: {att_err}")

            msg.close()

            full_text = chronological_body
            if attachment_texts:
                full_text += "\n\n" + "\n\n".join(attachment_texts)

            return subject, sender, msg_date, full_text, filenames

        except Exception as e:
            print(f"  └─► [MSG Parsing Error] Failed to parse '{file_path.name}': {e}")
            return f"Error reading {file_path.name}", "Unknown Sender", "Unknown Date", "", []


def fetch_extracted_email_payloads(account_info: dict, config: dict) -> List[Dict]:
    """
    Scans directory for ALL .eml and .msg files, cleans text/attachments,
    and returns 100% of parsed email payloads directly for LLM evaluation.
    """
    inbox_dir = Path(config.get("LOCAL_INBOX_DIR", "./outlook_inbox"))

    if not inbox_dir.exists():
        inbox_dir.mkdir(parents=True, exist_ok=True)
        print(f"[Outlook Extractor] Created folder at '{inbox_dir}'. Drop .eml or .msg files here.")
        return []

    sample_files = sorted(list(inbox_dir.glob("*.eml")) + list(inbox_dir.glob("*.msg")))

    if not sample_files:
        print(f"[Outlook Extractor] No .eml or .msg files found in '{inbox_dir}'.")
        return []

    print(f"\n=====================================================================")
    print(f" [Outlook Extractor] Ingesting ALL {len(sample_files)} sample email file(s) for LLM evaluation")
    print(f"=====================================================================")

    extracted_payloads = []

    for idx, file_path in enumerate(sample_files, start=1):
        if file_path.suffix.lower() == ".eml":
            subject, sender, msg_date, full_text, filenames = OutlookFileParser.parse_eml_file(file_path)
        elif file_path.suffix.lower() == ".msg":
            subject, sender, msg_date, full_text, filenames = OutlookFileParser.parse_msg_file(file_path)
        else:
            continue

        clean_subject = subject.lower().replace("re:", "").replace("fwd:", "").strip()
        parent_conv_id = f"THREAD-{abs(hash(clean_subject)) % 100000:05d}"

        extracted_payloads.append({
            "mailbox_account": account_info.get("account_name", "Outlook_Local_Desk"),
            "file_name": file_path.name,
            "email_sequence_index": idx,
            "parent_conversation_id": parent_conv_id,
            "sender": sender,
            "received_date": msg_date,
            "subject": subject,
            "combined_text_payload": full_text,
            "attached_files": filenames
        })

    return extracted_payloads